"""
Preprocessing Module for Resume and Candidate Screening System.
Handles multi-format text extraction (PDF, DOCX, TXT), text normalization,
PII anonymization, section segmentation, and entity extraction (experience, education).
"""

import re
import io
from typing import Dict, List, Optional, Tuple, Union
import nltk
from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer

# Initialize NLTK components safely
try:
    _STOPWORDS = set(stopwords.words('english'))
except Exception:
    nltk.download('stopwords', quiet=True)
    _STOPWORDS = set(stopwords.words('english'))

try:
    _LEMMATIZER = WordNetLemmatizer()
    _LEMMATIZER.lemmatize('testing')
except Exception:
    nltk.download('wordnet', quiet=True)
    _LEMMATIZER = WordNetLemmatizer()


# Common section header patterns in resumes
SECTION_PATTERNS = {
    'summary': re.compile(
        r'(?i)\b(professional summary|executive summary|summary|profile|about me|objective|career objective)\b'
    ),
    'skills': re.compile(
        r'(?i)\b(skills|technical skills|core competencies|technologies|expertise|tools & technologies|skill set)\b'
    ),
    'experience': re.compile(
        r'(?i)\b(work experience|professional experience|experience|employment history|work history|career history)\b'
    ),
    'education': re.compile(
        r'(?i)\b(education|academic background|academics|qualifications|degrees)\b'
    ),
    'projects': re.compile(
        r'(?i)\b(projects|key projects|academic projects|personal projects|portfolio)\b'
    ),
    'certifications': re.compile(
        r'(?i)\b(certifications|certificates|licenses|courses|accreditations)\b'
    )
}

# Degree level detection patterns
DEGREE_PATTERNS = [
    ('phd', re.compile(r'(?i)\b(ph\.?d|doctor of philosophy|doctorate)\b')),
    ('masters', re.compile(r'(?i)\b(master(?:\'s)?|m\.?s\.?|m\.?tech|m\.?sc|m\.?b\.?a|mca|m\.?eng)\b')),
    ('bachelors', re.compile(r'(?i)\b(bachelor(?:\'s)?|b\.?s\.?|b\.?tech|b\.?sc|b\.?e\.?|bca|bba|b\.?eng)\b')),
    ('diploma', re.compile(r'(?i)\b(diploma|associate(?:\'s)? degree|high school)\b'))
]

# Year of experience regex patterns
EXP_PATTERNS = [
    re.compile(r'(?i)(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?\s+(?:professional\s+)?experience'),
    re.compile(r'(?i)experience\s*(?:of)?\s*:?\s*(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)'),
    re.compile(r'(?i)over\s+(\d+(?:\.\d+)?)\s*(?:years?|yrs?)')
]


def clean_text(
    text: str,
    mask_pii: bool = False,
    remove_stopwords: bool = False,
    lemmatize: bool = False,
    lower: bool = True
) -> str:
    """
    Cleans and standardizes raw resume or job description text.

    Args:
        text: Raw input text.
        mask_pii: If True, replaces emails, phone numbers, and URLs with tokens
                  to support blind/anonymized recruitment screening.
        remove_stopwords: If True, filters out generic English stopwords.
        lemmatize: If True, applies WordNet morphological lemmatization.
        lower: If True, converts text to lowercase.

    Returns:
        Cleaned text string.
    """
    if not isinstance(text, str) or not text.strip():
        return ""

    # Normalize unicode quotation marks and dashes
    text = text.replace('\u2018', "'").replace('\u2019', "'")
    text = text.replace('\u201c', '"').replace('\u201d', '"')
    text = text.replace('\u2013', '-').replace('\u2014', '-')
    text = text.replace('\u2022', ' ').replace('\u25cf', ' ').replace('\u2219', ' ')

    # PII Masking or Removal
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
    url_pattern = r'https?://[^\s]+|www\.[^\s]+|linkedin\.com/in/[^\s]+|github\.com/[^\s]+'

    if mask_pii:
        text = re.sub(email_pattern, ' <EMAIL> ', text)
        text = re.sub(phone_pattern, ' <PHONE> ', text)
        text = re.sub(url_pattern, ' <URL> ', text)
    else:
        # Simple removal of URLs and raw emails for NLP vectorization
        text = re.sub(url_pattern, ' ', text)
        text = re.sub(email_pattern, ' ', text)
        text = re.sub(phone_pattern, ' ', text)

    if lower:
        text = text.lower()

    # Retain standard alphanumerics and essential symbols (+ for C++, # for C#, . for .NET, / - for terms)
    # Remove awkward punctuation
    text = re.sub(r'[^\w\s\+\#\.\/\-]', ' ', text)

    # Collapse repeated whitespace and newlines
    text = re.sub(r'\s+', ' ', text).strip()

    if remove_stopwords or lemmatize:
        words = text.split()
        processed_words = []
        for w in words:
            if remove_stopwords and w in _STOPWORDS:
                continue
            if lemmatize:
                w = _LEMMATIZER.lemmatize(w)
            processed_words.append(w)
        text = ' '.join(processed_words)

    return text


def extract_text_from_pdf(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a PDF file path, raw bytes, or BytesIO object.
    """
    import pypdf
    text = ""
    try:
        if isinstance(file_input, str):
            reader = pypdf.PdfReader(file_input)
        elif isinstance(file_input, bytes):
            reader = pypdf.PdfReader(io.BytesIO(file_input))
        else:
            reader = pypdf.PdfReader(file_input)

        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text.strip()


def extract_text_from_docx(file_input: Union[str, bytes, io.BytesIO]) -> str:
    """
    Extracts text from a DOCX file path, raw bytes, or BytesIO object.
    """
    import docx
    text = ""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_input))
        elif isinstance(file_input, io.BytesIO):
            doc = docx.Document(file_input)
        else:
            doc = docx.Document(file_input)

        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + " "
                text += "\n"
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return text.strip()


def extract_text_from_file(file_input: Union[str, bytes, io.BytesIO], filename: str = "") -> str:
    """
    Generic dispatcher to extract text from PDF, DOCX, or plain text formats.
    """
    lower_name = filename.lower()
    if lower_name.endswith('.pdf'):
        return extract_text_from_pdf(file_input)
    elif lower_name.endswith(('.docx', '.doc')):
        return extract_text_from_docx(file_input)
    else:
        # Fallback to UTF-8 / latin-1 decoding
        if isinstance(file_input, str):
            try:
                with open(file_input, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read().strip()
            except Exception:
                return ""
        elif isinstance(file_input, (bytes, bytearray)):
            try:
                return file_input.decode('utf-8', errors='ignore').strip()
            except Exception:
                return file_input.decode('latin-1', errors='ignore').strip()
        elif hasattr(file_input, 'read'):
            raw = file_input.read()
            if isinstance(raw, bytes):
                return raw.decode('utf-8', errors='ignore').strip()
            return str(raw).strip()
    return ""


def extract_sections(text: str) -> Dict[str, str]:
    """
    Segments resume text into structural sections:
    summary, skills, experience, education, projects, certifications.
    """
    lines = text.split('\n')
    sections: Dict[str, List[str]] = {
        'header': [],
        'summary': [],
        'skills': [],
        'experience': [],
        'education': [],
        'projects': [],
        'certifications': [],
        'other': []
    }

    current_section = 'header'

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        matched_section = None
        # Check if line looks like a section header (short line, match pattern)
        if len(stripped.split()) <= 5:
            for sec_name, pattern in SECTION_PATTERNS.items():
                if pattern.search(stripped):
                    matched_section = sec_name
                    break

        if matched_section:
            current_section = matched_section
        else:
            sections[current_section].append(stripped)

    return {k: '\n'.join(v).strip() for k, v in sections.items()}


def extract_years_of_experience(text: str) -> float:
    """
    Extracts explicit years of experience mentioned in resume/JD text.
    Returns estimated years (float). If no explicit number is found,
    estimates based on employment date ranges (e.g., 2019 - 2024).
    """
    # 1. Search for explicit mentions like '5+ years experience', '7 yrs of experience'
    found_years = []
    for pat in EXP_PATTERNS:
        matches = pat.findall(text)
        for m in matches:
            try:
                found_years.append(float(m))
            except ValueError:
                pass

    if found_years:
        return max(found_years)

    # 2. Date range heuristic: detect patterns like '2018 - 2023', '2019 to Present'
    date_pattern = re.compile(r'\b(19\d\d|20\d\d)\s*(?:-|–|—|to)\s*(19\d\d|20\d\d|present|current)\b', re.IGNORECASE)
    date_matches = date_pattern.findall(text)

    current_year = 2026  # Operational year reference
    total_span = 0.0

    for start_str, end_str in date_matches:
        try:
            start_yr = int(start_str)
            if end_str.lower() in ('present', 'current'):
                end_yr = current_year
            else:
                end_yr = int(end_str)

            if end_yr >= start_yr and (end_yr - start_yr) <= 35:
                total_span = max(total_span, float(end_yr - start_yr))
        except Exception:
            continue

    return total_span


def extract_education(text: str) -> Dict[str, Union[str, bool]]:
    """
    Identifies highest education degree level detected in the candidate text.
    """
    res = {
        'highest_degree': 'Not Specified',
        'has_phd': False,
        'has_masters': False,
        'has_bachelors': False,
        'has_diploma': False
    }

    for degree_label, pat in DEGREE_PATTERNS:
        if pat.search(text):
            if degree_label == 'phd':
                res['has_phd'] = True
                if res['highest_degree'] == 'Not Specified':
                    res['highest_degree'] = 'Ph.D.'
            elif degree_label == 'masters':
                res['has_masters'] = True
                if res['highest_degree'] in ('Not Specified', 'Bachelors', 'Diploma'):
                    res['highest_degree'] = 'Master\'s Degree'
            elif degree_label == 'bachelors':
                res['has_bachelors'] = True
                if res['highest_degree'] in ('Not Specified', 'Diploma'):
                    res['highest_degree'] = 'Bachelor\'s Degree'
            elif degree_label == 'diploma':
                res['has_diploma'] = True
                if res['highest_degree'] == 'Not Specified':
                    res['highest_degree'] = 'Diploma / Associate'

    return res
